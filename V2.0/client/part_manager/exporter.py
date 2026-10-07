"""Export part records to an Excel (.xlsx) file."""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from .models import Part

HEADER_FILL = PatternFill("solid", fgColor="3B5BDB")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")


def export_parts_to_excel(parts: list[Part], file_path: str,
                          project_name: str = "") -> str:
    """Write the given parts into an Excel workbook and return the path."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Parts"

    headers = ["Part Number", "Part Name", "Type", "Material",
               "Description", "Created"]
    ws.append(headers)
    for col, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER

    for part in parts:
        ws.append([part.part_number, part.part_name, part.part_type,
                   part.material, part.description, part.created_date])

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row,
                            min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center")

    widths = [16, 30, 12, 16, 36, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    ws.freeze_panes = "A2"
    wb.save(file_path)
    return file_path
