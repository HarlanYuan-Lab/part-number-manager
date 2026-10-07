"""Excel export service (creates a .xlsx workbook from the active database)."""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import Part

HEADERS = ["Part Number", "Part Name", "Part Type", "Material",
           "Description", "Created Date"]


def export_parts_to_excel(parts: list[Part], file_path: str) -> int:
    """Write the given parts into an .xlsx workbook; return the row count."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Parts"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="3B5BDB")
    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for r, part in enumerate(parts, start=2):
        values = [part.part_number, part.part_name, part.part_type,
                  part.material, part.description, part.created_date]
        for c, value in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=value).border = border

    for i, width in enumerate([16, 30, 14, 18, 40, 14], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(file_path)
    return len(parts)
